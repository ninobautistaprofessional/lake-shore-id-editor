# -*- coding: utf-8 -*-
"""Validate the generated deck: list each slide's key text."""
import sys
from pptx import Presentation
from pptx.util import Emu

# The deck uses box-drawing and bullet glyphs; force UTF-8 so printing them
# does not fail on a cp1252 Windows console.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

p = Presentation(r"presentation\Lake_Shore_ID_Management_System.pptx")
print("Slides:", len(p.slides))
for i, s in enumerate(p.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = " / ".join(x.strip() for x in sh.text_frame.text.splitlines() if x.strip())
            texts.append(t[:110])
        elif sh.has_table:
            texts.append("[TABLE " + str(len(sh.table.rows)) + "x" + str(len(sh.table.columns)) + "]")
        elif sh.shape_type == 13:
            texts.append("[PICTURE]")
    label = " | ".join(texts[:4])
    print(f"  {i:>2}. {label}")
print("DONE")